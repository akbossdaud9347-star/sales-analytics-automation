import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.chart import BarChart, PieChart, Reference

INPUT_FILE = "sales_data.xlsx"
OUTPUT_FILE = "automated_sales_report.xlsx"

print("Reading sales data...")

df = pd.read_excel(INPUT_FILE)

print("Records loaded:", len(df))

df["Date"] = pd.to_datetime(df["Date"])
df["Total Sales"] = df["Quantity"] * df["Unit Price"]
df["Month"] = df["Date"].dt.strftime("%B")
df["Year-Month"] = df["Date"].dt.strftime("%Y-%m")

total_sales = df["Total Sales"].sum()
total_orders = len(df)
total_quantity = df["Quantity"].sum()
average_order_value = total_sales / total_orders

monthly_sales = (
    df.groupby("Year-Month")["Total Sales"]
    .sum()
    .reset_index()
)

monthly_sales.columns = ["Month", "Total Sales"]

product_sales = (
    df.groupby("Product")
    .agg(
        Total_Quantity=("Quantity", "sum"),
        Total_Sales=("Total Sales", "sum")
    )
    .reset_index()
)

product_sales = product_sales.sort_values(
    "Total_Sales",
    ascending=False
)

category_sales = (
    df.groupby("Category")
    .agg(
        Total_Quantity=("Quantity", "sum"),
        Total_Sales=("Total Sales", "sum")
    )
    .reset_index()
)

salesperson_sales = (
    df.groupby("Salesperson")
    .agg(
        Total_Orders=("Customer ID", "count"),
        Total_Quantity=("Quantity", "sum"),
        Total_Sales=("Total Sales", "sum")
    )
    .reset_index()
)

salesperson_sales = salesperson_sales.sort_values(
    "Total_Sales",
    ascending=False
)

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="Sales Data",
        index=False
    )

    monthly_sales.to_excel(
        writer,
        sheet_name="Monthly Analysis",
        index=False
    )

    product_sales.to_excel(
        writer,
        sheet_name="Product Analysis",
        index=False
    )

    category_sales.to_excel(
        writer,
        sheet_name="Category Analysis",
        index=False
    )

    salesperson_sales.to_excel(
        writer,
        sheet_name="Salesperson Analysis",
        index=False
    )

wb = load_workbook(OUTPUT_FILE)

dashboard = wb.create_sheet("Dashboard", 0)

dashboard["A1"] = "AUTOMATED SALES ANALYTICS DASHBOARD"
dashboard["A1"].font = Font(
    bold=True,
    size=18
)

dashboard["A3"] = "Key Performance Indicators"

dashboard["A5"] = "Total Sales"
dashboard["B5"] = total_sales

dashboard["A6"] = "Total Orders"
dashboard["B6"] = total_orders

dashboard["A7"] = "Total Quantity Sold"
dashboard["B7"] = total_quantity

dashboard["A8"] = "Average Order Value"
dashboard["B8"] = average_order_value

for cell in ["A5", "A6", "A7", "A8"]:
    dashboard[cell].font = Font(bold=True)

for cell in ["B5", "B6", "B7", "B8"]:
    dashboard[cell].font = Font(
        bold=True,
        size=12
    )

dashboard["B5"].number_format = '#,##0.00'
dashboard["B8"].number_format = '#,##0.00'

monthly_sheet = wb["Monthly Analysis"]

bar_chart = BarChart()
bar_chart.title = "Monthly Sales"
bar_chart.y_axis.title = "Sales"
bar_chart.x_axis.title = "Month"

data = Reference(
    monthly_sheet,
    min_col=2,
    min_row=1,
    max_row=monthly_sheet.max_row
)

categories = Reference(
    monthly_sheet,
    min_col=1,
    min_row=2,
    max_row=monthly_sheet.max_row
)

bar_chart.add_data(
    data,
    titles_from_data=True
)

bar_chart.set_categories(categories)

bar_chart.height = 8
bar_chart.width = 14

dashboard.add_chart(
    bar_chart,
    "D4"
)

category_sheet = wb["Category Analysis"]

pie_chart = PieChart()

labels = Reference(
    category_sheet,
    min_col=1,
    min_row=2,
    max_row=category_sheet.max_row
)

data = Reference(
    category_sheet,
    min_col=3,
    min_row=1,
    max_row=category_sheet.max_row
)

pie_chart.add_data(
    data,
    titles_from_data=True
)

pie_chart.set_categories(labels)

pie_chart.title = "Sales by Category"

pie_chart.height = 8
pie_chart.width = 10

dashboard.add_chart(
    pie_chart,
    "D20"
)

for ws in wb.worksheets:

    ws.freeze_panes = "A2"

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center"
        )

    for column in ws.columns:
        column_letter = column[0].column_letter
        ws.column_dimensions[column_letter].width = 18

dashboard.column_dimensions["A"].width = 25
dashboard.column_dimensions["B"].width = 20

wb.save(OUTPUT_FILE)

print()
print("====================================")
print("AUTOMATED SALES REPORT CREATED")
print("====================================")
print("Total Sales:", round(total_sales, 2))
print("Total Orders:", total_orders)
print("Total Quantity:", total_quantity)
print("Average Order Value:", round(average_order_value, 2))
print("File:", OUTPUT_FILE)
print()
print("DONE!")